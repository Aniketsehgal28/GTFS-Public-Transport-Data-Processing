import os
import sys
import pandas as pd
import tkinter as tk
from tkinter import ttk, messagebox
from datetime import datetime
from collections import defaultdict
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import gtfs_kit as gk

if getattr(sys, 'frozen', False):
    BASE_DIR = sys._MEIPASS
else:
    BASE_DIR = os.path.dirname(os.path.abspath(__file__))


def load_gtfs(zip_file_path):
    try:
        feed = gk.read_feed(zip_file_path, dist_units="km")

        for required in ['routes', 'stops', 'trips', 'stop_times', 'calendar']:
            if getattr(feed, required) is None:
                raise ValueError(f"Missing required GTFS file: {required}.txt")

        trips_calendar = feed.trips.merge(feed.calendar, on='service_id')

        def get_service_flags(row):
            flags = []
            if any(row[day] for day in ['monday', 'tuesday', 'wednesday', 'thursday', 'friday']):
                flags.append('D')
            if row['saturday']:
                flags.append('Š')
            if row['sunday']:
                flags.append('S')
            return ''.join(flags)

        trips_calendar['service_flags'] = trips_calendar.apply(get_service_flags, axis=1)

        merged = (
            feed.stop_times
            .merge(trips_calendar[['trip_id', 'route_id', 'service_flags']], on='trip_id')
            .merge(feed.routes[['route_id', 'route_short_name']], on='route_id')
            .merge(feed.stops[['stop_id', 'stop_name']], on='stop_id')
        )

        stop_data = defaultdict(list)
        for (stop_name, route), group in merged.groupby(['stop_name', 'route_short_name']):
            time_flag_dict = defaultdict(set)
            for arrival in group['arrival_time']:
                try:
                    hour, minute, *_ = map(int, arrival.split(':'))
                    time = f"{hour:02}:{minute:02}"
                    for flag in group['service_flags'].iloc[0]:
                        time_flag_dict[time].add(flag)
                except:
                    continue
            stop_data[stop_name].append((route, time_flag_dict))

        return stop_data

    except Exception as e:
        print(f"Error loading GTFS data: {e}")
        return defaultdict(list)


def set_vertical_alignment(cell, alignment="center"):
    """Set vertical alignment for a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    vAlign = OxmlElement('w:vAlign')
    vAlign.set(qn('w:val'), alignment)
    tcPr.append(vAlign)


def export_schedule_to_docx(stop, schedule):
    filename = f"Sticker_{stop.replace(' ', '_')}.docx"
    doc = Document()

    # Title
    title_par = doc.add_paragraph()
    title_run = title_par.add_run(stop)
    title_run.bold = True
    title_run.font.size = Pt(36)
    title_run.font.name = 'Times New Roman'
    title_par.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Red date
    date_par = doc.add_paragraph()
    date_run = date_par.add_run(f"nuo {datetime.now().strftime('%Y %m %d')}")
    date_run.bold = True
    date_run.font.size = Pt(12)
    date_run.font.color.rgb = RGBColor(255, 0, 0)
    date_par.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()  # Spacer

    # Destination label
    dest_par = doc.add_paragraph()
    dest_par.add_run("|Vilniaus AS:").bold = True

    # Table
    table = doc.add_table(rows=0, cols=2)
    table.autofit = False
    table.columns[0].width = Inches(1.0)
    table.columns[1].width = Inches(5.0)

    for index, (route, time_flag_dict) in enumerate(schedule):
        row = table.add_row().cells

        # LEFT CELL (Route number)
        row[0].text = ""
        p = row[0].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(route.upper())
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.bold = True
        set_vertical_alignment(row[0], "center")

        # Background color
        shading_elm = OxmlElement('w:shd')
        shading_elm.set(qn('w:fill'), '001F5B')
        row[0]._tc.get_or_add_tcPr().append(shading_elm)

        # RIGHT CELL (Times)
        sorted_times = sorted(time_flag_dict.items())
        formatted = [f"{time} {''.join(sorted(flags))}" for time, flags in sorted_times]
        text_block = " ".join(formatted)
        right = row[1].paragraphs[0].add_run(text_block)
        right.font.size = Pt(10)
        row[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT

        # MINIMAL THIN SPACER between routes
        if index < len(schedule) - 1:
            spacer_row = table.add_row().cells
            for cell in spacer_row:
                paragraph = cell.paragraphs[0]
                run = paragraph.add_run(" ")
                run.font.size = Pt(1)
                paragraph.space_before = Pt(0)
                paragraph.space_after = Pt(0)

    doc.add_paragraph()  # Spacer

    # Footer line 1
    doc.add_paragraph("D - darbo dienomis, Š - šeštadieniais, S - sekmadieniais")

    # Footer line 2
    footer = doc.add_paragraph()
    footer.add_run("Aktualūs grafikai ir naujienos ").bold = True
    link = footer.add_run("WWW.VRAP.LT")
    link.bold = True
    link.font.color.rgb = RGBColor(0, 0, 255)
    link.font.underline = True
    footer.add_run(" arba TRAFI programėlėje")

    doc.save(filename)
    messagebox.showinfo("DOCX Exported", f"Saved as {filename}")


def format_time_suffixes(time_flag_dict):
    label_order = ['D', 'Š', 'S']
    return [f"{time} {''.join(flag for flag in label_order if flag in flags)}"
            for time, flags in sorted(time_flag_dict.items())]


def create_gui(stop_data):
    current_schedule = []

    def show_schedule():
        stop = stop_var.get()
        tree.delete(*tree.get_children())
        if stop not in stop_data:
            messagebox.showerror("Error", f"Stop '{stop}' not found.")
            return

        current_schedule.clear()
        for route, time_flag_dict in stop_data[stop]:
            formatted = format_time_suffixes(time_flag_dict)
            if formatted:
                current_schedule.append((route, time_flag_dict))
                tree.insert("", "end", values=(route.upper(), ', '.join(formatted)))

    def export_docx():
        if not current_schedule:
            messagebox.showerror("No Data", "Please generate the timetable before exporting.")
            return
        export_schedule_to_docx(stop_var.get(), current_schedule)

    root = tk.Tk()
    root.title("Vilnius Bus Timetable")
    root.geometry("800x550")
    root.configure(bg="#f8f9fa")

    ttk.Label(root, text="🚌 Vilnius Bus Timetable", font=("Segoe UI", 16, "bold")).pack(pady=10)

    top_frame = ttk.Frame(root)
    top_frame.pack(pady=5)

    ttk.Label(top_frame, text="Stop:").pack(side=tk.LEFT, padx=5)
    stop_var = tk.StringVar()
    stop_menu = ttk.Combobox(top_frame, textvariable=stop_var, values=sorted(stop_data), width=40, state="readonly")
    stop_menu.pack(side=tk.LEFT)

    if stop_data:
        stop_var.set(sorted(stop_data)[0])

    ttk.Button(top_frame, text="Show Timetable", command=show_schedule).pack(side=tk.LEFT, padx=10)
    ttk.Button(top_frame, text="Export Sticker (DOCX)", command=export_docx).pack(side=tk.LEFT)

    table_frame = ttk.Frame(root)
    table_frame.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)

    columns = ("Route", "Arrival Times")
    tree = ttk.Treeview(table_frame, columns=columns, show="headings", height=15)
    tree.heading("Route", text="Route")
    tree.heading("Arrival Times", text="Arrival Times")
    tree.column("Route", width=100, anchor=tk.CENTER)
    tree.column("Arrival Times", width=600, anchor=tk.W)

    vsb = ttk.Scrollbar(table_frame, orient="vertical", command=tree.yview)
    tree.configure(yscrollcommand=vsb.set)

    tree.grid(row=0, column=0, sticky="nsew")
    vsb.grid(row=0, column=1, sticky="ns")
    table_frame.grid_rowconfigure(0, weight=1)
    table_frame.grid_columnconfigure(0, weight=1)

    style = ttk.Style()
    style.configure("Treeview.Heading", font=("Segoe UI", 10, "bold"))
    style.configure("Treeview", font=("Segoe UI", 10), rowheight=25)

    root.mainloop()


if __name__ == "__main__":
    gtfs_path = os.path.join(BASE_DIR, "gtfs.zip")
    stop_data = load_gtfs(gtfs_path)
    create_gui(stop_data)
