import os
import zipfile
import pandas as pd
import tkinter as tk
from tkinter import ttk, messagebox
from datetime import datetime
from collections import defaultdict
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

GTFS_ZIP = "gtfs.zip"
EXTRACT_DIR = "gtfs_data"

def extract_gtfs():
    if not os.path.exists(EXTRACT_DIR):
        with zipfile.ZipFile(GTFS_ZIP, 'r') as zip_ref:
            zip_ref.extractall(EXTRACT_DIR)

def get_route_186_data():
    extract_gtfs()
    routes = pd.read_csv(f"{EXTRACT_DIR}/routes.txt")
    stops = pd.read_csv(f"{EXTRACT_DIR}/stops.txt")
    trips = pd.read_csv(f"{EXTRACT_DIR}/trips.txt")
    stop_times = pd.read_csv(f"{EXTRACT_DIR}/stop_times.txt")
    calendar = pd.read_csv(f"{EXTRACT_DIR}/calendar.txt")

    # Adjust this if route 186 is labeled differently
    route_186 = routes[routes['route_short_name'].astype(str).str.contains("186", case=False)]
    if route_186.empty:
        raise Exception("Route 186 not found.")

    route_ids = route_186['route_id'].unique()
    trips_186 = trips[trips['route_id'].isin(route_ids)]
    trip_ids = trips_186['trip_id'].unique()
    stop_times = stop_times[stop_times['trip_id'].isin(trip_ids)]

    merged = stop_times.merge(trips_186[['trip_id', 'route_id', 'service_id']], on='trip_id')
    merged = merged.merge(routes[['route_id', 'route_short_name']], on='route_id')
    merged = merged.merge(stops[['stop_id', 'stop_name']], on='stop_id')
    merged = merged.merge(calendar, on='service_id')

    def get_service_flags(row):
        flags = []
        if any(row[day] for day in ['monday', 'tuesday', 'wednesday', 'thursday', 'friday']):
            flags.append('D')
        if row['saturday']: flags.append('Š')
        if row['sunday']: flags.append('S')
        return ''.join(flags)

    merged['service_flags'] = merged.apply(get_service_flags, axis=1)

    stop_data = defaultdict(list)
    for (stop_name, route), group in merged.groupby(['stop_name', 'route_short_name']):
        time_flag_dict = defaultdict(set)
        for _, row in group.iterrows():
            try:
                hour, minute, *_ = map(int, row['arrival_time'].split(':'))
                time = f"{hour:02}:{minute:02}"
                for flag in row['service_flags']:
                    time_flag_dict[time].add(flag)
            except:
                continue
        stop_data[stop_name].append((route, time_flag_dict))
    return stop_data

def set_vertical_alignment(cell, alignment="center"):
    tcPr = cell._tc.get_or_add_tcPr()
    vAlign = OxmlElement('w:vAlign')
    vAlign.set(qn('w:val'), alignment)
    tcPr.append(vAlign)

def export_schedule_to_docx(stop, schedule):
    filename = f"Sticker_186_{stop.replace(' ', '_')}.docx"
    doc = Document()

    title = doc.add_paragraph()
    run = title.add_run(stop)
    run.bold = True
    run.font.size = Pt(36)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    date_par = doc.add_paragraph()
    date_run = date_par.add_run(f"nuo {datetime.now().strftime('%Y %m %d')}")
    date_run.bold = True
    date_run.font.size = Pt(12)
    date_run.font.color.rgb = RGBColor(255, 0, 0)
    date_par.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()
    doc.add_paragraph("|Vilniaus AS:").bold = True

    table = doc.add_table(rows=0, cols=2)
    table.autofit = False
    table.columns[0].width = Inches(1.0)
    table.columns[1].width = Inches(5.0)

    for index, (route, time_flag_dict) in enumerate(schedule):
        row = table.add_row().cells

        row[0].text = ""
        p = row[0].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(str(route).upper())
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.bold = True
        set_vertical_alignment(row[0], "center")

        shading = OxmlElement('w:shd')
        shading.set(qn('w:fill'), '001F5B')
        row[0]._tc.get_or_add_tcPr().append(shading)

        sorted_times = sorted(time_flag_dict.items())
        formatted = [f"{time} {''.join(sorted(flags))}" for time, flags in sorted_times]
        text_block = " ".join(formatted)
        right = row[1].paragraphs[0].add_run(text_block)
        right.font.size = Pt(10)
        row[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT

        if index < len(schedule) - 1:
            spacer = table.add_row().cells
            for cell in spacer:
                run = cell.paragraphs[0].add_run(" ")
                run.font.size = Pt(1)

    doc.add_paragraph()
    doc.add_paragraph("D - darbo dienomis, Š - šeštadieniais, S - sekmadieniais")
    footer = doc.add_paragraph()
    footer.add_run("Aktualūs grafikai ir naujienos ").bold = True
    link = footer.add_run("WWW.VRAP.LT")
    link.bold = True
    link.font.color.rgb = RGBColor(0, 0, 255)
    link.underline = True
    footer.add_run(" arba TRAFI programėlėje")

    doc.save(filename)
    messagebox.showinfo("DOCX Exported", f"Saved as {filename}")

def create_gui(stop_data):
    current_schedule = []

    def show_schedule():
        stop = stop_var.get()
        tree.delete(*tree.get_children())
        current_schedule.clear()
        if stop not in stop_data:
            messagebox.showerror("Error", f"Stop '{stop}' not found.")
            return

        for route, time_flag_dict in stop_data[stop]:
            formatted = [f"{time} {''.join(sorted(flags))}" for time, flags in sorted(time_flag_dict.items())]
            if formatted:
                current_schedule.append((route, time_flag_dict))
                tree.insert("", "end", values=(str(route).upper(), ', '.join(formatted)))

    def export_docx():
        if not current_schedule:
            messagebox.showerror("No Data", "Please generate the timetable before exporting.")
            return
        export_schedule_to_docx(stop_var.get(), current_schedule)

    root = tk.Tk()
    root.title("Route 186 Bus Timetable")
    root.geometry("800x550")
    root.configure(bg="#f8f9fa")

    ttk.Label(root, text="🚌 Route 186 Bus Timetable", font=("Segoe UI", 16, "bold")).pack(pady=10)

    top_frame = ttk.Frame(root)
    top_frame.pack(pady=5)

    ttk.Label(top_frame, text="Stop:").pack(side=tk.LEFT, padx=5)
    stop_var = tk.StringVar()
    stop_menu = ttk.Combobox(top_frame, textvariable=stop_var, values=sorted(stop_data), width=40, state="readonly")
    stop_menu.pack(side=tk.LEFT)

    if stop_data:
        stop_var.set(sorted(stop_data)[0])

    ttk.Button(top_frame, text="Show Timetable", command=show_schedule).pack(side=tk.LEFT, padx=10)
    ttk.Button(top_frame, text="Export Sticker (DOCX)", command=export_docx).pack(side=tk.LEFT)

    table_frame = ttk.Frame(root)
    table_frame.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)

    columns = ("Route", "Arrival Times")
    tree = ttk.Treeview(table_frame, columns=columns, show="headings", height=15)
    tree.heading("Route", text="Route")
    tree.heading("Arrival Times", text="Arrival Times")
    tree.column("Route", width=100, anchor=tk.CENTER)
    tree.column("Arrival Times", width=600, anchor=tk.W)

    vsb = ttk.Scrollbar(table_frame, orient="vertical", command=tree.yview)
    tree.configure(yscrollcommand=vsb.set)

    tree.grid(row=0, column=0, sticky="nsew")
    vsb.grid(row=0, column=1, sticky="ns")
    table_frame.grid_rowconfigure(0, weight=1)
    table_frame.grid_columnconfigure(0, weight=1)

    style = ttk.Style()
