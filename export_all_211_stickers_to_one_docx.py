import os
import zipfile
import pandas as pd
from collections import defaultdict
from datetime import datetime
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

GTFS_ZIP = "gtfs.zip"
EXTRACT_DIR = "gtfs_data"

def extract_gtfs():
    if not os.path.exists(EXTRACT_DIR):
        with zipfile.ZipFile(GTFS_ZIP, 'r') as zip_ref:
            zip_ref.extractall(EXTRACT_DIR)

def get_all_buses_at_stops_of_route(target_route_short_name):
    extract_gtfs()
    routes = pd.read_csv(f"{EXTRACT_DIR}/routes.txt")
    stops = pd.read_csv(f"{EXTRACT_DIR}/stops.txt")
    trips = pd.read_csv(f"{EXTRACT_DIR}/trips.txt")
    stop_times = pd.read_csv(f"{EXTRACT_DIR}/stop_times.txt")
    calendar = pd.read_csv(f"{EXTRACT_DIR}/calendar.txt")

    routes['route_short_name'] = routes['route_short_name'].astype(str)
    stops['stop_name'] = stops['stop_name'].astype(str)

    target_routes = routes[routes['route_short_name'] == target_route_short_name]
    if target_routes.empty:
        raise Exception(f"Route {target_route_short_name} not found in GTFS.")

    target_route_ids = target_routes['route_id'].unique()
    target_trip_ids = trips[trips['route_id'].isin(target_route_ids)]['trip_id'].unique()
    target_stops = stop_times[stop_times['trip_id'].isin(target_trip_ids)]['stop_id'].unique()

    trips_calendar = trips.merge(calendar, on='service_id')
    merged = (
        stop_times
        .merge(trips_calendar[['trip_id', 'route_id', 'service_id',
                               'monday', 'tuesday', 'wednesday', 'thursday', 'friday',
                               'saturday', 'sunday']], on='trip_id')
        .merge(routes[['route_id', 'route_short_name']], on='route_id')
        .merge(stops[['stop_id', 'stop_name']], on='stop_id')
    )

    def get_service_flags(row):
        flags = []
        if any(row[day] for day in ['monday', 'tuesday', 'wednesday', 'thursday', 'friday']):
            flags.append('D')
        if row['saturday']: flags.append('Š')
        if row['sunday']: flags.append('S')
        return ''.join(flags)

    merged['service_flags'] = merged.apply(get_service_flags, axis=1)

    stop_data = defaultdict(list)
    filtered = merged[merged['stop_id'].isin(target_stops)]
    for (stop_name, route), group in filtered.groupby(['stop_name', 'route_short_name']):
        time_flag_dict = defaultdict(set)
        for arrival in group['arrival_time']:
            try:
                hour, minute, *_ = map(int, arrival.split(':'))
                time = f"{hour:02}:{minute:02}"
                for flag in group['service_flags'].iloc[0]:
                    time_flag_dict[time].add(flag)
            except:
                continue
        stop_data[stop_name].append((route, time_flag_dict))

    return stop_data

def set_vertical_alignment(cell, alignment="center"):
    tcPr = cell._tc.get_or_add_tcPr()
    vAlign = OxmlElement('w:vAlign')
    vAlign.set(qn('w:val'), alignment)
    tcPr.append(vAlign)

def export_all_to_one_docx(stop_data):
    doc = Document()

    for stop, schedule in sorted(stop_data.items()):
        doc.add_page_break()
        title = doc.add_paragraph()
        run = title.add_run(stop)
        run.bold = True
        run.font.size = Pt(36)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        date_par = doc.add_paragraph()
        date_run = date_par.add_run(f"nuo {datetime.now().strftime('%Y %m %d')}")
        date_run.bold = True
        date_run.font.size = Pt(12)
        date_run.font.color.rgb = RGBColor(255, 0, 0)
        date_par.alignment = WD_ALIGN_PARAGRAPH.CENTER

        doc.add_paragraph()
        doc.add_paragraph("|Vilniaus AS:").bold = True

        table = doc.add_table(rows=0, cols=2)
        table.autofit = False
        table.columns[0].width = Inches(1.0)
        table.columns[1].width = Inches(5.0)

        for index, (route, time_flag_dict) in enumerate(schedule):
            row = table.add_row().cells

            row[0].text = ""
            p = row[0].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(str(route).upper())
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.bold = True
            set_vertical_alignment(row[0])

            shading = OxmlElement('w:shd')
            shading.set(qn('w:fill'), '001F5B')
            row[0]._tc.get_or_add_tcPr().append(shading)

            sorted_times = sorted(time_flag_dict.items())
            formatted = [f"{time} {''.join(sorted(flags))}" for time, flags in sorted_times]
            row[1].paragraphs[0].add_run(" ".join(formatted)).font.size = Pt(10)

        doc.add_paragraph()
        doc.add_paragraph("D - darbo dienomis, Š - šeštadieniais, S - sekmadieniais")
        footer = doc.add_paragraph()
        footer.add_run("Aktualūs grafikai ir naujienos ").bold = True
        link = footer.add_run("WWW.VRAP.LT")
        link.bold = True
        link.font.color.rgb = RGBColor(0, 0, 255)
        link.underline = True
        footer.add_run(" arba TRAFI programėlėje")

    doc.save("Route211_All_Stops_Combined.docx")
    print("✅ Exported all stickers into Route211_All_Stops_Combined.docx")

# Run it
stop_data = get_all_buses_at_stops_of_route("211")
export_all_to_one_docx(stop_data)