import os
import zipfile
import pandas as pd
from datetime import datetime
from collections import defaultdict
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

GTFS_ZIP = "gtfs.zip"
EXTRACT_DIR = "gtfs_data"

def extract_gtfs():
    if not os.path.exists(EXTRACT_DIR):
        with zipfile.ZipFile(GTFS_ZIP, 'r') as zip_ref:
            zip_ref.extractall(EXTRACT_DIR)

def load_data():
    extract_gtfs()
    routes = pd.read_csv(f"{EXTRACT_DIR}/routes.txt")
    stops = pd.read_csv(f"{EXTRACT_DIR}/stops.txt")
    trips = pd.read_csv(f"{EXTRACT_DIR}/trips.txt")
    stop_times = pd.read_csv(f"{EXTRACT_DIR}/stop_times.txt")
    calendar = pd.read_csv(f"{EXTRACT_DIR}/calendar.txt")

    stop_times_sorted = stop_times.sort_values(by=['trip_id', 'stop_sequence'])
    first_stops = stop_times_sorted.groupby('trip_id').first().reset_index()[['trip_id', 'stop_id']]
    last_stops = stop_times_sorted.groupby('trip_id').last().reset_index()[['trip_id', 'stop_id']]
    first_stops.columns = ['trip_id', 'first_stop_id']
    last_stops.columns = ['trip_id', 'last_stop_id']

    trips = trips.merge(first_stops, on='trip_id').merge(last_stops, on='trip_id')
    trips = trips.merge(stops[['stop_id', 'stop_name']], left_on='last_stop_id', right_on='stop_id', how='left')
    trips = trips.rename(columns={'stop_name': 'direction_stop_name'})

    return routes, stops, trips, stop_times, calendar

def get_schedule_by_direction(filter_by_route_prefix=None):
    routes, stops, trips, stop_times, calendar = load_data()

    routes['route_short_name'] = routes['route_short_name'].astype(str)
    stops['stop_name'] = stops['stop_name'].astype(str)

    if filter_by_route_prefix:
        filtered_routes = routes[routes['route_short_name'].str.startswith(filter_by_route_prefix)]
        route_ids = filtered_routes['route_id'].unique()
        trips = trips[trips['route_id'].isin(route_ids)]

    # First, enrich trips with required info
    enriched_trips = trips[['trip_id', 'route_id', 'service_id', 'direction_stop_name']]

    # Merge stop_times with enriched trips (brings route_id and service_id)
    stop_times = stop_times.merge(enriched_trips, on='trip_id', how='left')

    # Now safe to merge with calendar and routes
    stop_times = stop_times.merge(calendar, on='service_id', how='left')
    stop_times = stop_times.merge(routes[['route_id', 'route_short_name']], on='route_id', how='left')
    stop_times = stop_times.merge(stops[['stop_id', 'stop_name']], on='stop_id')

    def get_service_flags(row):
        flags = []
        if any(row[day] for day in ['monday', 'tuesday', 'wednesday', 'thursday', 'friday']):
            flags.append('D')
        if row['saturday']:
            flags.append('Š')
        if row['sunday']:
            flags.append('S')
        return ''.join(flags)

    stop_times['service_flags'] = stop_times.apply(get_service_flags, axis=1)

    stop_data = defaultdict(lambda: defaultdict(list))

    for ((stop_name, direction), group) in stop_times.groupby(['stop_name', 'direction_stop_name']):
        time_flag_dict_per_route = defaultdict(set)
        for _, row in group.iterrows():
            try:
                hour, minute, *_ = map(int, row['arrival_time'].split(':'))
                time = f"{hour:02}:{minute:02}"
                for flag in row['service_flags']:
                    time_flag_dict_per_route[(row['route_short_name'], time)].add(flag)
            except:
                continue
        organized = defaultdict(lambda: defaultdict(set))
        for (route, time), flags in time_flag_dict_per_route.items():
            organized[route][time] = flags
        for route, times in organized.items():
            stop_data[stop_name][direction].append((route, times))

    return stop_data

def export_schedule_to_docx(stop, direction, schedule):
    filename = f"Sticker_{stop.replace(' ', '_')}_to_{direction.replace(' ', '_')}.docx"
    doc = Document()

    # Title
    title = doc.add_paragraph()
    run = title.add_run(f"{stop} → {direction}")
    run.bold = True
    run.font.size = Pt(36)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Date
    date_par = doc.add_paragraph()
    date_run = date_par.add_run(f"nuo {datetime.now().strftime('%Y %m %d')}")
    date_run.bold = True
    date_run.font.size = Pt(12)
    date_run.font.color.rgb = RGBColor(255, 0, 0)
    date_par.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()
    doc.add_paragraph("|Vilniaus AS:")

    table = doc.add_table(rows=0, cols=2)
    table.autofit = False
    table.columns[0].width = Inches(0.6)
    table.columns[1].width = Inches(5.8)

    for route, time_flag_dict in schedule:
        row = table.add_row().cells
        cell = row[0]
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(str(route).upper())
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.bold = True
        set_vertical_alignment(cell)

        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), '001F5B')
        tcPr.append(shd)

        sorted_times = sorted(time_flag_dict.items())
        formatted = [f"{time} {''.join(sorted(flags))}" for time, flags in sorted_times]
        row[1].paragraphs[0].add_run(" ".join(formatted)).font.size = Pt(10)

    doc.add_paragraph()
    doc.add_paragraph("D - darbo dienomis, Š - šeštadieniais, S - sekmadieniais")

    footer = doc.add_paragraph()
    footer.add_run("Aktualūs grafikai ir naujienos ").bold = True
    link = footer.add_run("WWW.VRAP.LT")
    link.bold = True
    link.font.color.rgb = RGBColor(0, 0, 255)
    link.underline = True
    footer.add_run(" arba TRAFI programėlėje")

    doc.save(filename)
    print(f"✅ Saved {filename}")

def set_vertical_alignment(cell, alignment="center"):
    tcPr = cell._tc.get_or_add_tcPr()
    vAlign = OxmlElement('w:vAlign')
    vAlign.set(qn('w:val'), alignment)
    tcPr.append(vAlign)

# --- Run ---
directional_data = get_schedule_by_direction(filter_by_route_prefix="211")
for stop, directions in directional_data.items():
    for direction, sched in directions.items():
        export_schedule_to_docx(stop, direction, sched)
