import os
import zipfile
import pandas as pd
import tkinter as tk
from tkinter import ttk, messagebox
from datetime import datetime
from collections import defaultdict
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

GTFS_ZIP = "gtfs.zip"
EXTRACT_DIR = "gtfs_data"

def extract_gtfs():
    if not os.path.exists(EXTRACT_DIR):
        with zipfile.ZipFile(GTFS_ZIP, 'r') as zip_ref:
            zip_ref.extractall(EXTRACT_DIR)

def get_all_routes_for_211_stops():
    extract_gtfs()
    routes = pd.read_csv(f"{EXTRACT_DIR}/routes.txt")
    stops = pd.read_csv(f"{EXTRACT_DIR}/stops.txt")
    trips = pd.read_csv(f"{EXTRACT_DIR}/trips.txt")
    stop_times = pd.read_csv(f"{EXTRACT_DIR}/stop_times.txt")
    calendar = pd.read_csv(f"{EXTRACT_DIR}/calendar.txt")

    # Ensure route_short_name is string
    routes['route_short_name'] = routes['route_short_name'].astype(str)

    route_211 = routes[routes['route_short_name'].str.contains("211", case=False)]
    if route_211.empty:
        raise Exception("Route 211 not found.")

    route_211_ids = route_211['route_id'].unique()
    trip_ids_211 = trips[trips['route_id'].isin(route_211_ids)]['trip_id'].unique()

    stops_211 = stop_times[stop_times['trip_id'].isin(trip_ids_211)]
    stop_ids_211 = stops_211['stop_id'].unique()

    trips_calendar = trips.merge(calendar, on='service_id')
    merged = (
        stop_times
        .merge(trips_calendar[['trip_id', 'route_id', 'service_id', 'monday', 'tuesday', 'wednesday', 'thursday', 'friday', 'saturday', 'sunday']], on='trip_id')
        .merge(routes[['route_id', 'route_short_name']], on='route_id')
        .merge(stops[['stop_id', 'stop_name']], on='stop_id')
    )

    def get_service_flags(row):
        flags = []
        if any(row[day] for day in ['monday', 'tuesday', 'wednesday', 'thursday', 'friday']):
            flags.append('D')
        if row['saturday']:
            flags.append('Š')
        if row['sunday']:
            flags.append('S')
        return ''.join(flags)

    merged['service_flags'] = merged.apply(get_service_flags, axis=1)

    stop_data = defaultdict(list)
    for (stop_name, route), group in merged[merged['stop_id'].isin(stop_ids_211)].groupby(['stop_name', 'route_short_name']):
        time_flag_dict = defaultdict(set)
        for arrival in group['arrival_time']:
            try:
                hour, minute, *_ = map(int, arrival.split(':'))
                time = f"{hour:02}:{minute:02}"
                for flag in group['service_flags'].iloc[0]:
                    time_flag_dict[time].add(flag)
            except:
                continue
        stop_data[stop_name].append((route, time_flag_dict))
    return stop_data

def set_vertical_alignment(cell, alignment="center"):
    tcPr = cell._tc.get_or_add_tcPr()
    vAlign = OxmlElement('w:vAlign')
    vAlign.set(qn('w:val'), alignment)
    tcPr.append(vAlign)

def export_schedule_to_docx(stop, schedule):
    filename = f"Sticker_{stop.replace(' ', '_')}.docx"
    doc = Document()

    title = doc.add_paragraph()
    run = title.add_run(stop)
    run.bold = True
    run.font.size = Pt(36)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    date_par = doc.add_paragraph()
    date_run = date_par.add_run(f"nuo {datetime.now().strftime('%Y %m %d')}")
    date_run.bold = True
    date_run.font.size = Pt(12)
    date_run.font.color.rgb = RGBColor(255, 0, 0)
    date_par.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()
    doc.add_paragraph("|Vilniaus AS:").bold = True

    table = doc.add_table(rows=0, cols=2)
    table.autofit = False
    table.columns[0].width = Inches(1.0)
    table.columns[1].width = Inches(5.0)

    for index, (route, time_flag_dict) in enumerate(schedule):
        row = table.add_row().cells

        row[0].text = ""
        p = row[0].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(str(route).upper())
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.bold = True
        set_vertical_alignment(row[0], "center")

        shading = OxmlElement('w:shd')
        shading.set(qn('w:fill'), '001F5B')
        row[0]._tc.get_or_add_tcPr().append(shading)

        sorted_times = sorted(time_flag_dict.items())
        formatted = [f"{time} {''.join(sorted(flags))}" for time, flags in sorted_times]
        text_block = " ".join(formatted)
        right = row[1].paragraphs[0].add_run(text_block)
        right.font.size = Pt(10)
        row[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT

        if index < len(schedule) - 1:
            spacer = table.add_row().cells
            for cell in spacer:
                run = cell.paragraphs[0].add_run(" ")
                run.font.size = Pt(1)

    doc.add_paragraph()
    doc.add_paragraph("D - darbo dienomis, Š - šeštadieniais, S - sekmadieniais")
    footer = doc.add_paragraph()
    footer.add_run("Aktualūs grafikai ir naujienos ").bold = True
    link = footer.add_run("WWW.VRAP.LT")
    link.bold = True
    link.font.color.rgb = RGBColor(0, 0, 255)
    link.underline = True
    footer.add_run(" arba TRAFI programėlėje")

    doc.save(filename)
    messagebox.showinfo("DOCX Exported", f"Saved as {filename}")

def create_gui(stop_data):
    current_schedule = []

    def show_schedule():
        stop = stop_var.get()
        tree.delete(*tree.get_children())
        current_schedule.clear()
        if stop not in stop_data:
            messagebox.showerror("Error", f"Stop '{stop}' not found.")
            return

        for route, time_flag_dict in stop_data[stop]:
            formatted = [f"{time} {''.join(sorted(flags))}" for time, flags in sorted(time_flag_dict.items())]
            if formatted:
                current_schedule.append((route, time_flag_dict))
                tree.insert("", "end", values=(str(route).upper(), ', '.join(formatted)))

    def export_docx():
        if not current_schedule:
            messagebox.showerror("No Data", "Please generate the timetable before exporting.")
            return
        export_schedule_to_docx(stop_var.get(), current_schedule)

    root = tk.Tk()
    root.title("🚌 Bus Stop Timetable for Route 211")
    root.geometry("800x550")
    root.configure(bg="#f8f9fa")

    ttk.Label(root, text="Select a Stop on Route 211:", font=("Segoe UI", 14, "bold")).pack(pady=10)
    stop_var = tk.StringVar()
    stop_menu = ttk.Combobox(root, textvariable=stop_var, values=sorted(stop_data), width=50, state="readonly")
    stop_menu.pack()

    if stop_data:
        stop_var.set(sorted(stop_data)[0])

    ttk.Button(root, text="Show Timetable", command=show_schedule).pack(pady=5)
    ttk.Button(root, text="Export DOCX", command=export_docx).pack(pady=5)

    columns = ("Route", "Arrival Times")
    tree = ttk.Treeview(root, columns=columns, show="headings", height=15)
    tree.heading("Route", text="Route")
    tree.heading("Arrival Times", text="Arrival Times")
    tree.column("Route", width=100, anchor=tk.CENTER)
    tree.column("Arrival Times", width=600, anchor=tk.W)
    tree.pack(padx=10, pady=10, expand=True, fill=tk.BOTH)

    style = ttk.Style()
    style.configure("Treeview.Heading", font=("Segoe UI", 10, "bold"))
    style.configure("Treeview", font=("Segoe UI", 10), rowheight=25)

    root.mainloop()

if __name__ == "__main__":
    stop_data = get_all_routes_for_211_stops()
    create_gui(stop_data)