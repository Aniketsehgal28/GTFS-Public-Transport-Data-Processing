import os
import zipfile
import pandas as pd
import tkinter as tk
from tkinter import ttk, messagebox
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

GTFS_ZIP = "gtfs.zip"
EXTRACT_DIR = "gtfs_data"

def extract_and_load_gtfs():
    if not os.path.exists(EXTRACT_DIR):
        with zipfile.ZipFile(GTFS_ZIP, 'r') as zip_ref:
            zip_ref.extractall(EXTRACT_DIR)

    routes = pd.read_csv(f"{EXTRACT_DIR}/routes.txt")
    stops = pd.read_csv(f"{EXTRACT_DIR}/stops.txt")
    trips = pd.read_csv(f"{EXTRACT_DIR}/trips.txt")
    stop_times = pd.read_csv(f"{EXTRACT_DIR}/stop_times.txt")

    route_211 = routes[routes['route_short_name'] == 211]
    if route_211.empty:
        raise Exception("Route 211 not found.")

    route_ids = route_211['route_id'].unique()
    trip_ids = trips[trips['route_id'].isin(route_ids)]['trip_id'].unique()
    stop_seq = stop_times[stop_times['trip_id'].isin(trip_ids)]
    stop_seq = stop_seq.merge(stops[['stop_id', 'stop_name']], on='stop_id')

    # Choose first trip as example
    first_trip_id = stop_seq.sort_values(by=['trip_id', 'stop_sequence']).iloc[0]['trip_id']
    ordered = stop_seq[stop_seq['trip_id'] == first_trip_id].sort_values('stop_sequence')

    stop_to_routes = {}
    for stop_id, stop_name in ordered[['stop_id', 'stop_name']].drop_duplicates().values:
        related_trips = stop_times[stop_times['stop_id'] == stop_id]['trip_id'].unique()
        related_route_ids = trips[trips['trip_id'].isin(related_trips)]['route_id'].unique()
        route_names = routes[routes['route_id'].isin(related_route_ids)]['route_short_name'].unique().tolist()
        stop_to_routes[stop_name] = sorted(set(route_names))

    return stop_to_routes

def export_to_docx(stop, route_list):
    doc = Document()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(stop)
    run.font.size = Pt(28)
    run.bold = True

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Served by routes:")
    run.font.size = Pt(12)

    for route in route_list:
        p = doc.add_paragraph()
        run = p.add_run(f"• Route {route}")
        run.font.size = Pt(11)

    footer = doc.add_paragraph()
    footer.add_run("Data from VRAP / GTFS | Exported by your app").italic = True

    filename = f"Route_211_Stop_{stop.replace(' ', '_')}.docx"
    doc.save(filename)
    messagebox.showinfo("Exported", f"Saved: {filename}")

def create_gui(stop_data):
    root = tk.Tk()
    root.title("Route 211 - Stop Routes Viewer")
    root.geometry("600x400")

    ttk.Label(root, text="Select a Stop on Route 211:", font=("Segoe UI", 14)).pack(pady=10)
    stop_var = tk.StringVar()
    combo = ttk.Combobox(root, textvariable=stop_var, values=sorted(stop_data), width=50, state="readonly")
    combo.pack(pady=5)
    if stop_data:
        stop_var.set(sorted(stop_data)[0])

    result_text = tk.Text(root, height=10, width=70)
    result_text.pack(pady=10)

    def show_routes():
        stop = stop_var.get()
        routes = stop_data.get(stop, [])
        result_text.delete("1.0", tk.END)
        result_text.insert(tk.END, f"Routes serving {stop}:\n\n")
        for r in routes:
            result_text.insert(tk.END, f"• {r}\n")

    def export_doc():
        stop = stop_var.get()
        if stop:
            export_to_docx(stop, stop_data[stop])

    ttk.Button(root, text="Show Routes", command=show_routes).pack(pady=5)
    ttk.Button(root, text="Export to DOCX", command=export_doc).pack(pady=5)

    root.mainloop()

if __name__ == "__main__":
    try:
        data = extract_and_load_gtfs()
        create_gui(data)
    except Exception as e:
        print("Error:", e)
